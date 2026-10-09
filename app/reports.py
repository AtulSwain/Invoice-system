"""Monthly report for the CA: data, Excel export and one-page PDF export."""
import calendar
import io
from datetime import date

from .money import inr, rupees_decimal

AMOUNT_KEYS = ("total_qty", "net_total", "cgst", "sgst", "grand_total")


def month_bounds(year: int, month: int):
    last_day = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last_day)


def monthly_report(conn, year: int, month: int) -> dict:
    start, end = month_bounds(year, month)
    invoices = conn.execute(
        "SELECT * FROM invoices WHERE invoice_date BETWEEN ? AND ? ORDER BY invoice_date, source DESC, number",
        (start.isoformat(), end.isoformat()),
    ).fetchall()

    rows = []
    totals = {k: 0 for k in AMOUNT_KEYS}
    by_customer = {}
    for inv in invoices:
        cancelled = inv["status"] == "cancelled"
        amounts = {k: 0 if cancelled else inv[k] for k in AMOUNT_KEYS}
        status = "Cancelled" if cancelled else "Active"
        if inv["source"] == "imported":
            status += " (Imported)"
        rows.append({"id": inv["id"], "date": inv["invoice_date"], "number": inv["number"],
                     "source": inv["source"], "party": inv["customer_name"], "gstin": inv["customer_gstin"],
                     "status": status, "cancelled": cancelled, **amounts})
        for k in AMOUNT_KEYS:
            totals[k] += amounts[k]
        if not cancelled:
            key = inv["customer_name"].lower()
            cust = by_customer.setdefault(key, {"party": inv["customer_name"], "gstin": inv["customer_gstin"],
                                                "bills": 0, **{k: 0 for k in AMOUNT_KEYS}})
            cust["bills"] += 1
            for k in AMOUNT_KEYS:
                cust[k] += amounts[k]
            if not cust["gstin"] and inv["customer_gstin"]:
                cust["gstin"] = inv["customer_gstin"]

    live_numbers = sorted(r["number"] for r in rows if r["source"] == "live")
    imported_numbers = sorted(r["number"] for r in rows if r["source"] == "imported")
    missing = []
    if live_numbers:
        present = set(live_numbers)
        for n in range(live_numbers[0], live_numbers[-1] + 1):
            if n not in present:
                other = conn.execute(
                    "SELECT invoice_date FROM invoices WHERE source = 'live' AND number = ?", (n,)
                ).fetchone()
                missing.append({"number": n, "elsewhere": other["invoice_date"] if other else None})

    return {
        "year": year, "month": month, "month_name": f"{calendar.month_name[month]} {year}",
        "start": start, "end": end, "rows": rows, "totals": totals,
        "bill_count": len(rows),
        "cancelled_count": sum(1 for r in rows if r["cancelled"]),
        "live_count": len(live_numbers), "imported_count": len(imported_numbers),
        "first_number": live_numbers[0] if live_numbers else None,
        "last_number": live_numbers[-1] if live_numbers else None,
        "imported_first": imported_numbers[0] if imported_numbers else None,
        "imported_last": imported_numbers[-1] if imported_numbers else None,
        "missing": missing,
        "by_customer": sorted(by_customer.values(), key=lambda c: c["party"].lower()),
    }


def quick_check_text(rep) -> str:
    if not rep["rows"]:
        return f"No bills in {rep['month_name']}."
    parts = [f"{rep['bill_count']} bills"]
    if rep["first_number"] is not None:
        parts.append(f"invoice numbers {rep['first_number']} to {rep['last_number']}")
    if rep["imported_count"]:
        parts.append(f"{rep['imported_count']} imported (bill nos. {rep['imported_first']} to {rep['imported_last']})")
    if rep["cancelled_count"]:
        parts.append(f"{rep['cancelled_count']} cancelled")
    text = ", ".join(parts) + "."
    if rep["missing"]:
        notes = []
        for m in rep["missing"]:
            if m["elsewhere"]:
                y, mo, d = m["elsewhere"].split("-")
                notes.append(f"{m['number']} (dated {d}/{mo}/{y})")
            else:
                notes.append(f"{m['number']} (NOT FOUND)")
        text += " WARNING: numbers missing from this month: " + ", ".join(notes) + "."
    else:
        text += " No numbers missing."
    return text


def _dmy(iso: str) -> str:
    y, m, d = iso.split("-")
    return f"{d}/{m}/{y}"


# ---------------------------------------------------------------- Excel

def to_xlsx(rep, settings) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    money_fmt = "#,##,##0.00"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDDDDD")
    thin = Side(style="thin", color="999999")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)

    def sheet_header(ws, title, ncols):
        ws.append([f"{settings['company_name']}  |  GSTIN {settings['gstin']}"])
        ws.append([title])
        ws.append([quick_check_text(rep)])
        for r in (1, 2):
            ws.cell(row=r, column=1).font = Font(bold=True, size=13 if r == 1 else 12)
        ws.append([])
        return 5

    def table(ws, header_row, headings, widths, money_cols):
        for col, (h, w) in enumerate(zip(headings, widths), start=1):
            c = ws.cell(row=header_row, column=col, value=h)
            c.font, c.fill, c.border = bold, head_fill, box
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
            for c in row:
                c.border = box
                if c.column in money_cols:
                    c.number_format = money_fmt

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    hr = sheet_header(ws, f"Sales Register - {rep['month_name']}", 10)
    headings = ["Date", "Invoice No.", "Party Name", "Party GSTIN", "Total Quantity (Pcs)",
                "Taxable Value (Rs)", "CGST (Rs)", "SGST (Rs)", "Total (Rs)", "Status"]
    ws.append(headings)
    for r in rep["rows"]:
        ws.append([_dmy(r["date"]), r["number"], r["party"], r["gstin"], r["total_qty"],
                   rupees_decimal(r["net_total"]), rupees_decimal(r["cgst"]), rupees_decimal(r["sgst"]),
                   rupees_decimal(r["grand_total"]), r["status"]])
    t = rep["totals"]
    ws.append([f"TOTAL ({rep['bill_count']} bills)", "", "", "", t["total_qty"], rupees_decimal(t["net_total"]),
               rupees_decimal(t["cgst"]), rupees_decimal(t["sgst"]), rupees_decimal(t["grand_total"]), ""])
    for c in ws[ws.max_row]:
        c.font = bold
    table(ws, hr, headings, [12, 11, 30, 19, 12, 16, 13, 13, 16, 18], {6, 7, 8, 9})

    ws2 = wb.create_sheet("Summary by Customer")
    hr = sheet_header(ws2, f"Summary by Customer - {rep['month_name']}", 8)
    headings = ["Party Name", "Party GSTIN", "Bills", "Total Quantity (Pcs)", "Taxable Value (Rs)",
                "CGST (Rs)", "SGST (Rs)", "Total (Rs)"]
    ws2.append(headings)
    for c in rep["by_customer"]:
        ws2.append([c["party"], c["gstin"], c["bills"], c["total_qty"], rupees_decimal(c["net_total"]),
                    rupees_decimal(c["cgst"]), rupees_decimal(c["sgst"]), rupees_decimal(c["grand_total"])])
    ws2.append(["TOTAL", "", sum(c["bills"] for c in rep["by_customer"]), t["total_qty"],
                rupees_decimal(t["net_total"]), rupees_decimal(t["cgst"]), rupees_decimal(t["sgst"]),
                rupees_decimal(t["grand_total"])])
    for c in ws2[ws2.max_row]:
        c.font = bold
    table(ws2, hr, headings, [30, 19, 8, 12, 16, 13, 13, 16], {5, 6, 7, 8})

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- PDF (single page)

def to_pdf(rep, settings) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph, Table, TableStyle
    from reportlab.lib.styles import ParagraphStyle

    from .pdf import register_fonts

    fonts = register_fonts()
    page_w, page_h = landscape(A4)
    margin = 10 * mm
    t = rep["totals"]

    def money(p):
        return inr(p, symbol=False)

    inv_data = [["Date", "Invoice No.", "Party Name", "Party GSTIN", "Total Qty", "Taxable Value",
                 "CGST", "SGST", "Total", "Status"]]
    for r in rep["rows"]:
        inv_data.append([_dmy(r["date"]), str(r["number"]), r["party"], r["gstin"], str(r["total_qty"]),
                         money(r["net_total"]), money(r["cgst"]), money(r["sgst"]), money(r["grand_total"]),
                         r["status"]])
    inv_data.append([f"TOTAL ({rep['bill_count']} bills)", "", "", "", str(t["total_qty"]), money(t["net_total"]),
                     money(t["cgst"]), money(t["sgst"]), money(t["grand_total"]), ""])

    cust_data = [["Party Name", "Bills", "Total Qty", "Taxable Value", "CGST", "SGST", "Total"]]
    for c in rep["by_customer"]:
        cust_data.append([c["party"], str(c["bills"]), str(c["total_qty"]), money(c["net_total"]),
                          money(c["cgst"]), money(c["sgst"]), money(c["grand_total"])])
    cust_data.append(["TOTAL", str(sum(c["bills"] for c in rep["by_customer"])), str(t["total_qty"]),
                      money(t["net_total"]), money(t["cgst"]), money(t["sgst"]), money(t["grand_total"])])

    def style(n_rows, right_from):
        return TableStyle([
            ("FONT", (0, 0), (-1, -1), fonts["sans"], 8),
            ("FONT", (0, 0), (-1, 0), fonts["sans_bold"], 8),
            ("FONT", (0, n_rows - 1), (-1, n_rows - 1), fonts["sans_bold"], 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDDDDD")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("ALIGN", (right_from, 0), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
        ])

    inv_table = Table(inv_data, repeatRows=1)
    inv_table.setStyle(style(len(inv_data), 4))
    cust_table = Table(cust_data)
    cust_table.setStyle(style(len(cust_data), 1))

    head_style = ParagraphStyle("h", fontName=fonts["sans_bold"], fontSize=13, leading=16)
    sub_style = ParagraphStyle("s", fontName=fonts["sans"], fontSize=9, leading=11)
    warn = bool(rep["missing"])
    blocks = [
        Paragraph(f"{settings['company_name']} &nbsp; | &nbsp; GSTIN {settings['gstin']}", head_style),
        Paragraph(f"Monthly Sales Report: <b>{rep['month_name']}</b>", head_style),
        Paragraph(("<font color='#B00000'>" if warn else "") + _esc(quick_check_text(rep)) +
                  ("</font>" if warn else ""), sub_style),
        inv_table,
        Paragraph("Summary by Customer", ParagraphStyle("h2", parent=head_style, fontSize=11, spaceBefore=6)),
        cust_table,
    ]
    gap = 4 * mm
    avail_w = page_w - 2 * margin
    sizes = [b.wrap(avail_w, 10_000) for b in blocks]
    total_h = sum(h for _, h in sizes) + gap * (len(blocks) - 1)
    max_w = max(w for w, _ in sizes)
    # Shrink everything to fit on a single page when the month is busy
    scale = min(1.0, (page_h - 2 * margin) / total_h, avail_w / max_w)

    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=landscape(A4))
    c.setTitle(f"Sales report {rep['month_name']}")
    c.translate(margin, page_h - margin)
    c.scale(scale, scale)
    y = 0
    for block, (w, h) in zip(blocks, sizes):
        y -= h
        block.drawOn(c, 0, y)
        y -= gap
    c.showPage()
    c.save()
    return out.getvalue()


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
