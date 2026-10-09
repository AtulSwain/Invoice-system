import io
from datetime import date

import pytest

from app import reports
from app.db import today_ist
from app.invoices import InvoiceError, cancel_invoice, create_invoice
from tests.conftest import invoice_data


def make(conn, owner, name="Jainam Creation", lines=((10, "125.50"),), inv_date=None):
    return create_invoice(conn, owner, invoice_data(name, lines, inv_date))


def test_cancel_keeps_number_and_blocks_changes(conn, owner, confirmed):
    a = make(conn, owner)
    b = make(conn, owner)
    cancel_invoice(conn, owner, a, "wrong rate")
    inv = conn.execute("SELECT * FROM invoices WHERE id = ?", (a,)).fetchone()
    assert inv["status"] == "cancelled" and inv["number"] == 198
    # number not reused
    c = make(conn, owner)
    assert conn.execute("SELECT number FROM invoices WHERE id = ?", (c,)).fetchone()[0] == 200
    # cannot cancel twice, cannot edit
    with pytest.raises(InvoiceError):
        cancel_invoice(conn, owner, a, "again")
    with pytest.raises(Exception, match="cannot be changed"):
        conn.execute("UPDATE invoices SET status = 'active' WHERE id = ?", (a,))
    with pytest.raises(Exception, match="cannot be changed"):
        conn.execute("UPDATE invoices SET net_total = 1 WHERE id = ?", (b,))
    with pytest.raises(Exception, match="cannot be changed"):
        conn.execute("UPDATE invoice_items SET qty = 1 WHERE invoice_id = ?", (b,))


def test_cancel_needs_reason(conn, owner, confirmed):
    a = make(conn, owner)
    with pytest.raises(InvoiceError):
        cancel_invoice(conn, owner, a, "  ")


def test_report_totals_equal_sum_of_invoices(conn, owner, confirmed):
    t = today_ist()
    ids = [make(conn, owner, "Jainam Creation", ((10, "125.50"), (3, "99.99"))),
           make(conn, owner, "Swastik Enterprises", ((7, "333.33"),)),
           make(conn, owner, "jainam creation", ((1, "1.01"),)),  # same customer, different case
           make(conn, owner, "Hareesh Enterprises", ((50, "80"),))]
    cancel_invoice(conn, owner, ids[3], "duplicate")
    rep = reports.monthly_report(conn, t.year, t.month)
    active = conn.execute("SELECT * FROM invoices WHERE status = 'active'").fetchall()
    for key in ("net_total", "cgst", "sgst", "grand_total", "total_qty"):
        assert rep["totals"][key] == sum(i[key] for i in active)
        assert sum(r[key] for r in rep["rows"]) == rep["totals"][key]
        assert sum(c[key] for c in rep["by_customer"]) == rep["totals"][key]
    cancelled_row = next(r for r in rep["rows"] if r["id"] == ids[3])
    assert cancelled_row["status"] == "Cancelled"
    assert all(cancelled_row[k] == 0 for k in reports.AMOUNT_KEYS)
    assert rep["bill_count"] == 4 and rep["cancelled_count"] == 1
    assert (rep["first_number"], rep["last_number"]) == (198, 201)
    assert rep["missing"] == []
    assert len(rep["by_customer"]) == 2  # Jainam merged, Hareesh cancelled
    assert "No numbers missing" in reports.quick_check_text(rep)


def test_report_warns_about_number_dated_in_another_month(conn, owner, confirmed):
    t = today_ist()
    first_of_month = date(t.year, t.month, 1)
    fy_start = date(t.year if t.month >= 4 else t.year - 1, 4, 1)
    if first_of_month == fy_start:
        pytest.skip("cannot back-date into a previous month in April")
    make(conn, owner)
    prev = date(first_of_month.year if first_of_month.month > 1 else first_of_month.year - 1,
                first_of_month.month - 1 if first_of_month.month > 1 else 12, 15)
    make(conn, owner, inv_date=prev)  # number 199, dated last month
    make(conn, owner)
    rep = reports.monthly_report(conn, t.year, t.month)
    assert [m["number"] for m in rep["missing"]] == [199]
    assert "WARNING" in reports.quick_check_text(rep)


def test_excel_and_pdf_exports(conn, owner, confirmed):
    from openpyxl import load_workbook
    t = today_ist()
    make(conn, owner)
    rep = reports.monthly_report(conn, t.year, t.month)
    from app.db import get_settings
    wb = load_workbook(io.BytesIO(reports.to_xlsx(rep, get_settings(conn))))
    assert wb.sheetnames == ["Invoices", "Summary by Customer"]
    headers = [c.value for c in wb["Invoices"][5]]
    assert headers[:4] == ["Date", "Invoice No.", "Party Name", "Party GSTIN"]
    pdf = reports.to_pdf(rep, get_settings(conn))
    assert pdf.startswith(b"%PDF") and pdf.count(b"/Type /Page\n") + pdf.count(b"/Type /Page ") >= 1
    from reportlab.lib.utils import isBytes  # noqa: F401  (reportlab present)
    assert pdf.count(b"/Type /Pages") == 1


def test_busy_month_pdf_is_one_page(conn, owner, confirmed):
    from app.db import get_settings
    t = today_ist()
    for _ in range(120):
        make(conn, owner)
    pdf = reports.to_pdf(reports.monthly_report(conn, t.year, t.month), get_settings(conn))
    assert b"/Count 1" in pdf
